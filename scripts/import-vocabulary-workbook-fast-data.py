#!/usr/bin/env python3
"""Import the supplied vocabulary workbook into the app's local lookup indexes."""

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = Path.home() / "Downloads" / "词汇总表20240307_new_带例句（含四级六级）.xlsx"
OUTPUT_DIR = ROOT / "src" / "data" / "vocabulary"
DEFINITION_CORRECTIONS = json.loads((OUTPUT_DIR / "definition-corrections.json").read_text(encoding="utf-8"))
DEFINITION_SHEET = "总24555"
EXAMPLE_PREFIX = re.compile(r"^\s*例句\s*[：:]\s*", re.I)
TRANSLATION_PREFIX = re.compile(r"^\s*翻译\s*[：:]\s*", re.I)
DEFINITION_PREFIX = re.compile(r"^\s*释义\s*[：:]\s*", re.I)
CHINESE = re.compile(r"[\u3400-\u9fff]")
WORD_SHAPE = re.compile(r"^[A-Za-z][A-Za-z'’ -]{0,78}$")

IRREGULAR_FORMS = {
    "be": "am is are was were been being", "go": "goes went gone", "do": "does did done",
    "have": "has had", "say": "says said", "make": "makes made", "take": "takes took taken",
    "come": "comes came", "get": "gets got gotten", "see": "sees saw seen", "give": "gives gave given",
    "find": "finds found", "think": "thinks thought", "buy": "buys bought", "bring": "brings brought",
    "build": "builds built", "feel": "feels felt", "leave": "leaves left", "meet": "meets met",
    "pay": "pays paid", "run": "runs ran", "sit": "sits sat", "speak": "speaks spoke spoken",
    "write": "writes wrote written", "drive": "drives drove driven", "eat": "eats ate eaten",
    "know": "knows knew known", "read": "reads", "sleep": "sleeps slept", "keep": "keeps kept",
    "hold": "holds held", "tell": "tells told", "understand": "understands understood",
    "hear": "hears heard", "lose": "loses lost", "send": "sends sent", "stand": "stands stood",
    "win": "wins won", "wear": "wears wore worn", "begin": "begins began begun", "break": "breaks broke broken",
    "choose": "chooses chose chosen", "fall": "falls fell fallen", "fly": "flies flew flown",
    "grow": "grows grew grown", "throw": "throws threw thrown", "draw": "draws drew drawn",
    "blow": "blows blew blown", "catch": "catches caught", "teach": "teaches taught", "sell": "sells sold",
    "spend": "spends spent", "mean": "means meant", "lead": "leads led", "set": "sets",
    "cut": "cuts", "hit": "hits", "put": "puts", "hurt": "hurts",
}


def cell_text(value):
    return value.strip() if isinstance(value, str) else ""


def normalize_word(value):
    return value.lower().strip().strip("-'’ ")


def word_forms(word):
    word = normalize_word(word)
    forms = {word}
    if word in IRREGULAR_FORMS:
        forms.update(IRREGULAR_FORMS[word].split())
    if " " in word:
        return forms
    if word.endswith("y") and len(word) > 1:
        forms.update({word[:-1] + "ies", word[:-1] + "ied", word[:-1] + "ier", word[:-1] + "iest"})
    if word.endswith("e"):
        forms.update({word + "d", word[:-1] + "ing", word + "s", word + "r", word + "st"})
    else:
        forms.update({word + "s", word + "es", word + "ed", word + "ing", word + "er", word + "est"})
        if len(word) > 2 and word[-1] not in "aeiouwxy" and word[-2] in "aeiou" and word[-3] not in "aeiou":
            forms.update({word + word[-1] + suffix for suffix in ("ed", "ing", "er", "est")})
    if word.endswith(("s", "x", "z", "ch", "sh", "o")):
        forms.add(word + "es")
    return forms


def contains_word(example, forms):
    for form in forms:
        if re.search(r"(?<![A-Za-z])" + re.escape(form) + r"(?![A-Za-z])", example, re.I):
            return True
    return False


def clean_example(value):
    value = cell_text(value)
    return EXAMPLE_PREFIX.sub("", value).strip()


def clean_translation(value):
    value = cell_text(value)
    return TRANSLATION_PREFIX.sub("", value).strip()


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w", encoding="utf-8") as handle:
        json.dump(data, handle, ensure_ascii=False, separators=(",", ":"))
        handle.write("\n")
    tmp.replace(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", nargs="?", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    if not args.workbook.is_file():
        raise SystemExit(f"Workbook not found: {args.workbook}")

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    definitions = {}
    definition_rows = 0
    for row in workbook[DEFINITION_SHEET].iter_rows(values_only=True):
        cells = list(row)
        if len(cells) < 5:
            continue
        level, raw_word, phonetic, definition_cn, raw_english = map(cell_text, cells[:5])
        word = normalize_word(raw_word)
        english_definition = DEFINITION_PREFIX.sub("", raw_english).strip()
        if not WORD_SHAPE.fullmatch(raw_word) or not word or not (definition_cn or english_definition):
            continue
        record = definitions.setdefault(word, {"definitionCn": [], "englishDefinitions": [], "level": "", "phonetic": ""})
        for value in (definition_cn,):
            if value and value not in record["definitionCn"]:
                record["definitionCn"].append(value)
        if english_definition and english_definition not in record["englishDefinitions"]:
            record["englishDefinitions"].append(english_definition)
        if not record["level"]:
            record["level"] = level
        if not record["phonetic"]:
            record["phonetic"] = phonetic
        definition_rows += 1

    details = {
        word: {
            "definitionCn": DEFINITION_CORRECTIONS.get(word, {}).get("definitionCn", "\n".join(record["definitionCn"])),
            "englishDefinitions": DEFINITION_CORRECTIONS.get(word, {}).get("englishDefinitions", record["englishDefinitions"]),
            "level": record["level"],
            "phonetic": DEFINITION_CORRECTIONS.get(word, {}).get("phonetic", record["phonetic"]),
        }
        for word, record in definitions.items()
    }

    examples_by_word = defaultdict(list)
    example_rows = 0
    for sheet in workbook.worksheets:
        if sheet.title == DEFINITION_SHEET:
            continue
        for row in sheet.iter_rows(values_only=True):
            cells = list(row)
            if len(cells) < 6:
                continue
            raw_word = cell_text(cells[1])
            word = normalize_word(raw_word)
            if not WORD_SHAPE.fullmatch(raw_word) or not word:
                continue
            forms = word_forms(word)
            for index in range(4, len(cells)):
                raw_example = cell_text(cells[index])
                example = clean_example(raw_example)
                if len(example) < 12 or CHINESE.search(example) or not re.search(r"[A-Za-z]", example) or not contains_word(example, forms):
                    continue
                translation = ""
                for possible_translation in cells[index + 1:]:
                    text = clean_translation(possible_translation)
                    if text and CHINESE.search(text):
                        translation = text
                        break
                if not translation:
                    continue
                examples_by_word[word].append((example, translation))
                example_rows += 1
                break

    examples_path = args.output_dir / "workbook-usage-examples.json"
    existing_examples = json.loads(examples_path.read_text(encoding="utf-8")) if examples_path.exists() else {}
    merged_examples = {normalize_word(key): value for key, value in existing_examples.items()}
    for word, candidates in examples_by_word.items():
        example, translation = min(candidates, key=lambda item: (len(item[0]) > 140, len(item[0]), len(item[1])))
        merged_examples[word] = {"englishText": example, "chineseText": translation}

    write_json(args.output_dir / "workbook-vocabulary-details.json", details)
    write_json(examples_path, merged_examples)
    workbook.close()

    print(f"Definition rows imported: {definition_rows}; unique words: {len(details)}")
    print(f"Verified bilingual workbook examples imported: {len(examples_by_word)} from {example_rows} candidate rows")
    print(f"Local example index now contains: {len(merged_examples)} words")


if __name__ == "__main__":
    main()
